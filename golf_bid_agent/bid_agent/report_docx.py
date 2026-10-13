"""Client-facing bid proposal (.docx)."""
from __future__ import annotations

from datetime import date, timedelta

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

GREEN = RGBColor(0x2E, 0x5E, 0x3E)
GREY = RGBColor(0x55, 0x55, 0x55)
RED = RGBColor(0xC0, 0x00, 0x00)


def _shade(cell, hex_fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_fill)
    tcPr.append(shd)


def _bottom_rule(paragraph, color="2E5E3E"):
    pPr = paragraph._p.get_or_add_pPr()
    bdr = OxmlElement("w:pBdr")
    b = OxmlElement("w:bottom")
    for k, v in (("w:val", "single"), ("w:sz", "8"), ("w:space", "4"), ("w:color", color)):
        b.set(qn(k), v)
    bdr.append(b)
    pPr.append(bdr)


def _money(x: float) -> str:
    return f"${x:,.2f}"


def _scope_bullets(est: dict) -> list[str]:
    p = est["project"]
    mats = {m["code"]: m for m in est["materials"]}
    out = []
    if p["earthwork_cy"]:
        out.append(f"Mobilize, protect the work area, and perform approximately {p['earthwork_cy']:,.0f} CY of "
                   "excavation, rough grading and earthwork.")
    if "MAT-001" in mats:
        label = "green" if "green" in p["project_type"] else "tee"
        out.append(f"Construct approximately {p['green_area_sqft']:,.0f} SF of {label} surface with a "
                   f"{p['rootzone_depth_in']:g}-in USGA rootzone mix (about {mats['MAT-001']['order_qty']:,.0f} CY "
                   "delivered), including finish shaping.")
    gb = [m for m in est["materials"] if "gravel blanket" in m["item"]]
    if gb:
        out.append(f"Install a {est['assumptions']['material_factors']['green_gravel_blanket_depth_in']}-in "
                   "drainage gravel blanket beneath the rootzone.")
    if p["drainage_pipe_lf"]:
        out.append(f"Install approximately {p['drainage_pipe_lf']:,.0f} LF of 4-in perforated HDPE drainage pipe "
                   "with stone backfill" + (" and geotextile trench wrap" if "MAT-012" in mats else "") + ".")
    if p["irrigation_pipe_lf"]:
        heads = mats.get("MAT-006", {}).get("order_qty")
        out.append(f"Install approximately {p['irrigation_pipe_lf']:,.0f} LF of PVC irrigation pipe"
                   + (" (4-in mainline and 2-in laterals)" if "MAT-004" in mats else "")
                   + (f", approximately {heads:,.0f} sprinkler heads" if heads else "")
                   + ", fittings and connections to the existing system.")
    if p["bunker_area_sqft"]:
        out.append(f"Rebuild approximately {p['bunker_area_sqft']:,.0f} SF of bunkers including shaping, liner, "
                   "drainage and washed bunker sand.")
    if p["sod_area_sqft"]:
        turf = next((m["item"] for m in est["materials"] if m["category"] == "turf"), "turf")
        out.append(f"Supply and install approximately {p['sod_area_sqft']:,.0f} SF of {turf}, "
                   "including tie-ins to existing turf.")
    for x in est["overrides"].get("extra_lines") or []:
        out.append(x["description"] + ".")
    out.append("Final cleanup and demobilization.")
    return out


def build_proposal(est: dict, path: str, cfg: dict, approval: dict | None = None) -> str:
    approved = bool(approval)
    co = cfg["company"]
    p, s = est["project"], est["summary"]
    today = date.fromisoformat(est["estimate_date"])
    valid_until = today + timedelta(days=int(co.get("bid_validity_days", 30)))

    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Inches(8.5), Inches(11)
    for side in ("left_margin", "right_margin"):
        setattr(sec, side, Inches(0.9))
    sec.top_margin = sec.bottom_margin = Inches(0.8)
    st = doc.styles["Normal"]
    st.font.name, st.font.size = "Arial", Pt(10)
    st.element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")
    for name in ("Heading 1", "Heading 2"):
        h = doc.styles[name]
        h.font.name, h.font.color.rgb = "Arial", GREEN
        h.font.size = Pt(13 if name == "Heading 1" else 11)
        h.element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")

    # header / footer
    hp = sec.header.paragraphs[0]
    hp.text = ""
    r = hp.add_run(co["name"])
    r.bold, r.font.size, r.font.color.rgb = True, Pt(11), GREEN
    hp.add_run(f"\t\t{co['address']}  |  {co['phone']}  |  {co['email']}").font.size = Pt(8)
    if not approved:
        dp = sec.header.add_paragraph()
        dr = dp.add_run("DRAFT - FOR INTERNAL REVIEW - NOT FOR RELEASE")
        dr.bold, dr.font.size, dr.font.color.rgb = True, Pt(9), RED
        dp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    fp = sec.footer.paragraphs[0]
    fp.text = f"Proposal BID-{p['project_id']}-{today:%Y%m%d}  |  {p['project_name']}"
    fp.runs[0].font.size, fp.runs[0].font.color.rgb = Pt(8), GREY

    t = doc.add_paragraph()
    tr = t.add_run("Bid Proposal")
    tr.bold, tr.font.size, tr.font.color.rgb = True, Pt(22), GREEN
    sub = doc.add_paragraph()
    sr = sub.add_run(p["project_name"])
    sr.font.size, sr.font.color.rgb = Pt(13), GREY
    _bottom_rule(sub)

    info = [("Proposal no.", f"BID-{p['project_id']}-{today:%Y%m%d}"), ("Date", f"{today:%B %d, %Y}"),
            ("Project ID", p["project_id"]), ("Location", p["location"]),
            ("Prepared by", f"{co['estimator']}, {co['name']}"), ("Valid until", f"{valid_until:%B %d, %Y}")]
    tbl = doc.add_table(rows=3, cols=4)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, (k, v) in enumerate(info):
        row, col = i % 3, (i // 3) * 2
        kc, vc = tbl.cell(row, col), tbl.cell(row, col + 1)
        kc.text, vc.text = k, str(v)
        kc.paragraphs[0].runs[0].bold = True
        kc.width = vc.width = Inches(1.2 if col == 0 else 1.2)
    for row in tbl.rows:
        row.cells[1].width = row.cells[3].width = Inches(2.15)

    doc.add_heading("1. Project understanding", level=1)
    doc.add_paragraph(
        f"{co['name']} is pleased to submit this proposal for the {p['project_type'].replace('_', ' ')} at "
        f"{p['project_name']}. Our price is based on the quantities in your request "
        f"and {p['site_access']} site access. Request notes: {p['notes']}")

    doc.add_heading("2. Scope of work", level=1)
    for b in _scope_bullets(est):
        doc.add_paragraph(b, style="List Bullet")

    doc.add_heading("3. Pricing", level=1)
    sov = est["schedule_of_values"]
    direct = sum(sov.values()) or 1
    mark = s["contingency"] + s["markup"]
    rows = [(k, round(v + v / direct * mark, 2)) for k, v in sov.items() if v]
    diff = round(s["total_bid"] - sum(a for _, a in rows), 2)  # rounding remainder to largest line
    if rows and diff:
        i = max(range(len(rows)), key=lambda j: rows[j][1])
        rows[i] = (rows[i][0], round(rows[i][1] + diff, 2))
    pt = doc.add_table(rows=1, cols=2)
    pt.style = "Table Grid"
    hdr = pt.rows[0].cells
    hdr[0].text, hdr[1].text = "Work category", "Amount"
    for c in hdr:
        _shade(c, "2E5E3E")
        c.paragraphs[0].runs[0].bold = True
        c.paragraphs[0].runs[0].font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
    hdr[1].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT
    for k, v in rows:
        rc = pt.add_row().cells
        rc[0].text, rc[1].text = k, _money(v)
        rc[1].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT
    tc = pt.add_row().cells
    tc[0].text, tc[1].text = "TOTAL LUMP SUM BID", _money(s["total_bid"])
    for c in tc:
        _shade(c, "E8F0E9")
        c.paragraphs[0].runs[0].bold = True
    tc[1].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT
    for row in pt.rows:
        row.cells[0].width, row.cells[1].width = Inches(4.9), Inches(1.8)
    doc.add_paragraph(
        "Prices include materials, labor, equipment, supervision, overhead, profit and an allowance for "
        "unforeseen conditions. Line amounts include a proportional share of those allowances.").runs[0].font.size = Pt(8)

    doc.add_heading("4. Schedule", level=1)
    sc = est["schedule"]
    win = sc["suggested_completion_window_days"]
    doc.add_paragraph(
        f"Estimated duration: {sc['target_duration_days']} working days"
        + (f" (allow up to {win} working days for weather and site conditions)" if win > sc["target_duration_days"]
           else "")
        + f" from mobilization, with a crew of approximately {est['crew']['crew_size']} including a full-time foreman. "
        "Start date to be agreed upon award.")

    doc.add_heading("5. Assumptions and clarifications", level=1)
    assumptions = [
        "Quantities are based on the dimensions provided in the project request and are subject to field "
        "verification. Changes in quantity will be adjusted at the unit rates underlying this proposal.",
        f"Site access assumed to be {p['site_access']}, with a staging area provided by the Owner near the work.",
    ]
    turf = next((m["item"] for m in est["materials"] if m["category"] == "turf"), None)
    if turf:
        assumptions.append(f"Turf is priced as {turf}. Other varieties can be priced on request.")
    if p["irrigation_pipe_lf"]:
        assumptions.append("Irrigation pipe sizing and head spacing are assumed pending the irrigation design; "
                           "connection to an existing, functioning water supply and control system.")
    if p["rootzone_depth_in"] and "MAT-001" in {m["code"] for m in est["materials"]}:
        assumptions.append("Rootzone mix assumed to meet USGA recommendations; laboratory testing by Owner.")
    for a in assumptions:
        doc.add_paragraph(a, style="List Bullet")
    qs = [f["client_question"] for f in est["flags"] if f["level"] == "CLARIFY" and f.get("client_question")]
    if qs:
        doc.add_heading("Information requested before award", level=2)
        for q in qs:
            doc.add_paragraph(q, style="List Bullet")

    doc.add_heading("6. Exclusions", level=1)
    for ex in est["exclusions"]:
        doc.add_paragraph(ex, style="List Bullet")

    doc.add_heading("7. Terms", level=1)
    doc.add_paragraph(f"This proposal is valid until {valid_until:%B %d, %Y}. Payment terms: [to be confirmed]. "
                      "Work will be performed in a workmanlike manner according to standard practice.")

    doc.add_heading("Acceptance", level=1)
    doc.add_paragraph("The above prices, scope and conditions are accepted. You are authorized to perform the work "
                      "as specified.")
    at = doc.add_table(rows=3, cols=2)
    labels = [("Owner / Club", co["name"]), ("Signature: ____________________________",
                                             "Signature: ____________________________"),
              ("Name / Date: __________________________", "Name / Date: __________________________")]
    for i, (a, b) in enumerate(labels):
        at.cell(i, 0).text, at.cell(i, 1).text = a, b
    at.cell(0, 0).paragraphs[0].runs[0].bold = at.cell(0, 1).paragraphs[0].runs[0].bold = True

    if approved:
        ap = doc.add_paragraph()
        ar = ap.add_run(f"Internal: approved by {approval.get('approved_by')} on {approval.get('approved_at')}.")
        ar.font.size, ar.font.color.rgb = Pt(7), GREY
    doc.save(path)
    return path
