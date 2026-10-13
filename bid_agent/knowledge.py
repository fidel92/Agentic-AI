"""Local retrieval (RAG) over your company documents and the workbook's text.

Everything runs on your machine: documents are split into short passages and ranked
with BM25 keyword scoring (no embeddings service, no extra API key, nothing uploaded).
Only the few passages a search returns are sent to Claude.

Sources indexed
  - <workspace>/knowledge/  : .md .txt .docx .pdf .csv files you add (specs, past
                              proposals, company policies, supplier quotes, field notes...)
  - the bid workbook        : Estimation_Rules, Historical_Projects notes, Project_Requests notes
"""
from __future__ import annotations

import csv
import math
import re
from collections import Counter
from dataclasses import dataclass, field
from pathlib import Path

from . import data as _data

SUPPORTED = {".md", ".txt", ".docx", ".pdf", ".csv"}
CHUNK_WORDS = 160
OVERLAP_WORDS = 40
MAX_FILE_MB = 25

STOP = set("""a an and are as at be by for from has have in is it its of on or that the this to was were
will with not no but if then than so such into over under per via can may should must shall our your
their they we you he she them his her i me my""".split())


def _tokens(text: str) -> list[str]:
    out = []
    for w in re.findall(r"[a-z0-9]+(?:[-.][a-z0-9]+)*", text.lower()):
        if w in STOP or len(w) < 2:
            continue
        if len(w) > 4 and w.endswith("ies"):
            w = w[:-3] + "y"
        elif len(w) > 3 and w.endswith("s") and not w.endswith("ss"):
            w = w[:-1]
        out.append(w)
    return out


@dataclass
class Passage:
    source: str            # file name or workbook sheet
    location: str          # page / paragraph / row reference
    kind: str              # "document" | "rules" | "history" | "requests"
    text: str
    tokens: list[str] = field(default_factory=list, repr=False)


# ------------------------------------------------------------------ readers
def _read_text(p: Path) -> list[tuple[str, str]]:
    return [("", p.read_text(encoding="utf-8", errors="ignore"))]


def _read_docx(p: Path) -> list[tuple[str, str]]:
    from docx import Document
    d = Document(str(p))
    parts = [para.text for para in d.paragraphs if para.text.strip()]
    for t in d.tables:
        for row in t.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                parts.append(" | ".join(cells))
    return [("", "\n\n".join(parts))]


def _read_pdf(p: Path) -> list[tuple[str, str]]:
    try:
        from pypdf import PdfReader
    except ImportError:
        raise RuntimeError("PDF support needs:  pip install pypdf")
    r = PdfReader(str(p))
    return [(f"page {i + 1}", page.extract_text() or "") for i, page in enumerate(r.pages)]


def _read_csv(p: Path) -> list[tuple[str, str]]:
    with open(p, newline="", encoding="utf-8", errors="ignore") as f:
        rows = list(csv.DictReader(f))
    return [(f"row {i + 2}", "; ".join(f"{k}: {v}" for k, v in r.items() if v)) for i, r in enumerate(rows)]


READERS = {".md": _read_text, ".txt": _read_text, ".docx": _read_docx, ".pdf": _read_pdf, ".csv": _read_csv}


def _chunk(text: str) -> list[str]:
    """Split on paragraphs, then pack into ~CHUNK_WORDS-word passages with overlap."""
    words: list[str] = []
    chunks: list[str] = []
    for para in re.split(r"\n\s*\n", text):
        pw = para.split()
        if not pw:
            continue
        if words and len(words) + len(pw) > CHUNK_WORDS:
            chunks.append(" ".join(words))
            words = words[-OVERLAP_WORDS:]
        words.extend(pw)
        while len(words) > CHUNK_WORDS * 1.5:            # very long paragraph
            chunks.append(" ".join(words[:CHUNK_WORDS]))
            words = words[CHUNK_WORDS - OVERLAP_WORDS:]
    if words:
        chunks.append(" ".join(words))
    return chunks


# -------------------------------------------------------------------- index
class KnowledgeBase:
    def __init__(self, bid_data=None):
        self.bid_data = bid_data
        self.passages: list[Passage] = []
        self.warnings: list[str] = []
        self._signature = None
        self.folder = None

    # -- building
    def _folder(self) -> Path:
        return _data.resolve("knowledge")

    def _scan(self) -> list[Path]:
        folder = self._folder()
        if not folder.is_dir():
            return []
        files = []
        for p in sorted(folder.rglob("*")):
            if p.is_file() and p.suffix.lower() in SUPPORTED and not p.name.startswith(("~$", ".")) \
                    and p.name.lower() != "readme.txt":
                try:
                    _data.resolve(p)                   # folder lock: refuses symlinks pointing outside
                except PermissionError as e:
                    self.warnings.append(str(e))
                    continue
                files.append(p)
        return files

    def refresh(self, force: bool = False) -> "KnowledgeBase":
        self.warnings = []
        files = self._scan()
        sig = tuple((str(f), f.stat().st_mtime, f.stat().st_size) for f in files) + (id(self.bid_data),)
        if not force and sig == self._signature:
            return self
        self.folder = self._folder()
        self.passages = []
        for f in files:
            if f.stat().st_size > MAX_FILE_MB * 1_000_000:
                self.warnings.append(f"{f.name}: skipped (over {MAX_FILE_MB} MB)")
                continue
            try:
                sections = READERS[f.suffix.lower()](f)
            except Exception as e:  # unreadable file should not stop the agent
                self.warnings.append(f"{f.name}: {e}")
                continue
            rel = str(f.relative_to(self.folder))
            if f.suffix.lower() == ".csv":
                for loc, txt in sections:
                    self._add(rel, loc, "document", txt)
                continue
            for loc, txt in sections:
                for i, ch in enumerate(_chunk(txt), 1):
                    self._add(rel, f"{loc + ', ' if loc else ''}passage {i}", "document", ch)
        self._add_workbook()
        self._build_stats()
        self._signature = sig
        return self

    def _add(self, source, location, kind, text):
        text = re.sub(r"\s+", " ", text).strip()
        if text:
            self.passages.append(Passage(source, location, kind, text))

    def _add_workbook(self):
        d = self.bid_data
        if d is None:
            return
        wb = d.source.name
        for _, r in d.rules.iterrows():
            self._add(f"{wb} / Estimation_Rules", r["rule_id"], "rules",
                      f"{r['rule_id']} {r['rule_name']}: {r['formula_or_logic']} -> {r['output_unit']}. {r['notes']}")
        for _, r in d.history.iterrows():
            self._add(f"{wb} / Historical_Projects", r["history_id"], "history",
                      f"{r['history_id']} {r['project_type']} in {r['region']}, {r['site_access']} access: "
                      f"final cost ${r['final_cost_usd']:,.0f}, planned {r['planned_duration_days']} days, actual "
                      f"{r['actual_duration_days']} days, crew {r['average_crew_size']}, weather delay "
                      f"{r['weather_delay_days']} days, data quality {r['data_quality']}. Notes: {r['notes']}")
        for _, r in d.projects.iterrows():
            self._add(f"{wb} / Project_Requests", r["project_id"], "requests",
                      f"{r['project_id']} {r['project_name']} ({r['project_type']}, {r['location']}, "
                      f"{r['site_access']} access). Notes: {r['notes']}")

    def _build_stats(self):
        for p in self.passages:
            p.tokens = _tokens(f"{p.source} {p.text}")
        self.N = len(self.passages)
        self.avgdl = (sum(len(p.tokens) for p in self.passages) / self.N) if self.N else 1
        df = Counter()
        for p in self.passages:
            df.update(set(p.tokens))
        self.idf = {t: math.log(1 + (self.N - n + 0.5) / (n + 0.5)) for t, n in df.items()}
        self.tf = [Counter(p.tokens) for p in self.passages]

    # -- searching
    def search(self, query: str, top_k: int = 5, scope: str = "all", k1: float = 1.5, b: float = 0.75) -> list[dict]:
        self.refresh()
        q = _tokens(query)
        if not q or not self.passages:
            return []
        kinds = {"all": None, "documents": {"document"}, "rules": {"rules"}, "history": {"history"},
                 "workbook": {"rules", "history", "requests"}}.get(scope)
        scored = []
        for i, p in enumerate(self.passages):
            if kinds and p.kind not in kinds:
                continue
            tf, dl, s = self.tf[i], len(p.tokens), 0.0
            for t in q:
                if t in tf:
                    f = tf[t]
                    s += self.idf.get(t, 0) * f * (k1 + 1) / (f + k1 * (1 - b + b * dl / self.avgdl))
            if s > 0:
                scored.append((s, i))
        scored.sort(reverse=True)
        out = []
        for s, i in scored[:max(1, min(top_k, 15))]:
            p = self.passages[i]
            out.append({"source": p.source, "location": p.location, "type": p.kind, "score": round(s, 2),
                        "text": p.text})
        return out

    def stats(self) -> dict:
        self.refresh()
        c = Counter(p.kind for p in self.passages)
        files = sorted({p.source for p in self.passages if p.kind == "document"})
        return {"folder": str(self.folder), "documents": files, "passages": dict(c), "warnings": self.warnings}
