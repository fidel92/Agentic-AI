"""Internal working estimate workbook (live formulas; blue = editable input)."""
from __future__ import annotations

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

FONT = "Arial"
F = Font(name=FONT, size=10)
FB = Font(name=FONT, size=10, bold=True)
F_IN = Font(name=FONT, size=10, color="0000FF")
F_LINK = Font(name=FONT, size=10, color="008000")
F_H1 = Font(name=FONT, size=14, bold=True, color="1F3B2C")
F_HDR = Font(name=FONT, size=10, bold=True, color="FFFFFF")
FILL_HDR = PatternFill("solid", fgColor="2E5E3E")
FILL_KEY = PatternFill("solid", fgColor="FFFF00")
FILL_TOT = PatternFill("solid", fgColor="E8F0E9")
FILL = {"CLARIFY": PatternFill("solid", fgColor="F8D7DA"), "REVIEW": PatternFill("solid", fgColor="FFF3CD"),
        "INFO": PatternFill("solid", fgColor="E7F1FA")}
THIN = Side(style="thin", color="BBBBBB")
BOX = Border(top=THIN, bottom=THIN, left=THIN, right=THIN)
USD = '$#,##0.00;($#,##0.00);"-"'
USD0 = '$#,##0;($#,##0);"-"'
QTY = '#,##0.##;(#,##0.##);"-"'
PCT = '0.0%;(0.0%);"-"'
WRAP = Alignment(wrap_text=True, vertical="top")


def _hdr(ws, row, headers, widths=None):
    for i, h in enumerate(headers, 1):
        c = ws.cell(row=row, column=i, value=h)
        c.font, c.fill, c.border = F_HDR, FILL_HDR, BOX
        c.alignment = Alignment(wrap_text=True, vertical="center")
    if widths:
        for i, w in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = ws.cell(row=row + 1, column=1)


def _cell(ws, row, col, value, font=F, fmt=None, fill=None, border=True, align=None):
    c = ws.cell(row=row, column=col, value=value)
    c.font = font
    if fmt:
        c.number_format = fmt
    if fill:
        c.fill = fill
    if border:
        c.border = BOX
    if align:
        c.alignment = align
    return c


def _title(ws, est, subtitle, approved):
    p = est["project"]
    ws["A1"] = f"{p['project_id']} - {p['project_name']}"
    ws["A1"].font = F_H1
    ws["A2"] = subtitle + ("" if approved else "   |   DRAFT - FOR INTERNAL REVIEW, NOT FOR RELEASE")
    ws["A2"].font = Font(name=FONT, size=10, italic=True, color="555555" if approved else "C00000")


def build_workbook(est: dict, path: str, cfg: dict, approval: dict | None = None) -> str:
    approved = bool(approval)
    wb = Workbook()
    p, s = est["project"], est["summary"]

    # ------------------------------------------------------------ Materials
    wm = wb.active
    wm.title = "Materials"
    _title(wm, est, "Materials - R-001 / R-002 / R-005", approved)
    hdr = ["Code", "Item", "Unit", "Base qty", "Waste %", "Round up to", "Order qty", "Unit cost ($)",
           "Extended ($)", "Schedule of values", "Basis / assumption"]
    _hdr(wm, 4, hdr, [9, 36, 9, 11, 9, 9, 11, 12, 14, 26, 70])
    r = 5
    for m in est["materials"]:
        _cell(wm, r, 1, m["code"])
        _cell(wm, r, 2, m["item"])
        _cell(wm, r, 3, m["unit"])
        _cell(wm, r, 4, m["base_qty"], F_IN, QTY)
        _cell(wm, r, 5, m["waste_factor"], F_IN, PCT)
        _cell(wm, r, 6, m["round_step"], F_IN, "0")
        _cell(wm, r, 7, f"=CEILING(ROUND(D{r}*(1+E{r}),6),F{r})", F, QTY)
        _cell(wm, r, 8, m["unit_cost"], F_IN, USD)
        _cell(wm, r, 9, f"=ROUND(G{r}*H{r},2)", F, USD)
        _cell(wm, r, 10, m["sov"])
        _cell(wm, r, 11, m["basis"], align=WRAP)
        r += 1
    mat_first, mat_last = 5, max(r - 1, 5)
    _cell(wm, r, 8, "Materials total", FB, fill=FILL_TOT)
    _cell(wm, r, 9, f"=SUM(I{mat_first}:I{mat_last})", FB, USD, FILL_TOT)
    mat_total_ref = f"Materials!$I${r}"
    wm.cell(row=r + 2, column=1, value="Blue = editable input. Order qty = base x (1 + waste), rounded up "
            "(R-002). Rootzone rounds up to 5 CY.").font = Font(name=FONT, size=9, italic=True)

    # ---------------------------------------------------------------- Labor
    wl = wb.create_sheet("Labor")
    _title(wl, est, "Labor & crew - R-003 / R-004 / R-011", approved)
    c = est["crew"]
    crew_inputs = [("Duration (working days)", c["duration_days"], "0"),
                   ("Crew size (incl. foreman)", c["crew_size"], "0"),
                   ("Foremen on site", c["foremen"], "0"),
                   ("Paid hours per person-day", c["paid_hours_per_day"], "0.0"),
                   ("Productive hours per person-day", c["productive_hours_per_day"], "0.0")]
    for i, (lab, val, fmt) in enumerate(crew_inputs, start=4):
        _cell(wl, i, 1, lab, FB)
        _cell(wl, i, 2, val, F_IN, fmt, FILL_KEY)
    _cell(wl, 9, 1, "Workers required (task h / prod h / days)", FB)
    lab_first = 12
    n_tasks = sum(1 for l in est["labor"] if l["productivity"])
    task_last = lab_first + max(n_tasks, 1) - 1
    _cell(wl, 9, 2, f"=ROUNDUP(SUM(G{lab_first}:G{task_last})/(B8*B4),0)", F, "0")
    wl.cell(row=4, column=3, value=f"Project-type crew range {c['type_min_crew']}-{c['type_max_crew']}; "
            f"comparables median {c['comps_median_crew']}. Crew = max(required + foreman, type minimum).").font = \
        Font(name=FONT, size=9, italic=True)
    hdr = ["Code", "Role", "Quantity", "Qty unit", "Productivity (per h)", "Access factor", "Hours",
           "Loaded rate ($/h)", "Cost ($)", "Schedule of values", "Basis"]
    _hdr(wl, 11, hdr, [10, 32, 11, 9, 13, 10, 10, 12, 13, 26, 70])
    wl.freeze_panes = "A12"
    r = lab_first
    for l in est["labor"]:
        _cell(wl, r, 1, l["code"])
        _cell(wl, r, 2, l["role"])
        _cell(wl, r, 4, l["qty_unit"])
        if l["productivity"]:
            _cell(wl, r, 3, l["quantity"], F_IN, QTY)
            _cell(wl, r, 5, l["productivity"], F_IN, QTY)
            _cell(wl, r, 6, l["access_factor"], F_IN, "0.00")
            _cell(wl, r, 7, f"=ROUND(C{r}/E{r}*F{r},2)", F, QTY)
        elif l["code"] == "LAB-008":
            _cell(wl, r, 3, "=B4", F, "0")
            _cell(wl, r, 7, "=B6*B4*B7", F, QTY)
        elif l["code"] == "SUPPORT":
            _cell(wl, r, 3, "=B5-B6", F, "0")
            _cell(wl, r, 7, f"=ROUND(MAX(0,(B5-B6)*B4*B7-SUM(G{lab_first}:G{task_last})),2)", F, QTY)
        else:
            _cell(wl, r, 3, l["quantity"], F_IN, QTY)
            _cell(wl, r, 7, l["hours"], F_IN, QTY)
        for col in (5, 6):
            if not wl.cell(row=r, column=col).value:
                _cell(wl, r, col, None)
        _cell(wl, r, 8, l["rate"], F_IN, USD)
        _cell(wl, r, 9, f"=ROUND(G{r}*H{r},2)", F, USD)
        _cell(wl, r, 10, l["sov"])
        _cell(wl, r, 11, l["basis"], align=WRAP)
        r += 1
    _cell(wl, r, 6, "Totals", FB, fill=FILL_TOT)
    _cell(wl, r, 7, f"=SUM(G{lab_first}:G{r - 1})", FB, QTY, FILL_TOT)
    _cell(wl, r, 8, None, fill=FILL_TOT)
    _cell(wl, r, 9, f"=SUM(I{lab_first}:I{r - 1})", FB, USD, FILL_TOT)
    lab_last, lab_total_ref = r - 1, f"Labor!$I${r}"
    if not any(l["code"] == "SUPPORT" for l in est["labor"]):
        wl.cell(row=r + 2, column=1, value="Crew support line omitted (task hours exceed paid crew time or "
                "excluded by reviewer).").font = Font(name=FONT, size=9, italic=True)

    # ------------------------------------------------------------ Equipment
    we = wb.create_sheet("Equipment")
    _title(we, est, "Equipment rental", approved)
    _hdr(we, 4, ["Code", "Equipment", "Typical use", "Days", "Daily rate ($)", "Cost ($)", "Basis"],
         [9, 22, 20, 8, 13, 13, 80])
    r = 5
    for e in est["equipment"]:
        _cell(we, r, 1, e["code"])
        _cell(we, r, 2, e["name"])
        _cell(we, r, 3, e["use"])
        _cell(we, r, 4, e["days"], F_IN, "0")
        _cell(we, r, 5, e["rate"], F_IN, USD)
        _cell(we, r, 6, f"=ROUND(D{r}*E{r},2)", F, USD)
        _cell(we, r, 7, e["basis"], align=WRAP)
        r += 1
    _cell(we, r, 5, "Equipment total", FB, fill=FILL_TOT)
    _cell(we, r, 6, f"=SUM(F5:F{max(r - 1, 5)})", FB, USD, FILL_TOT)
    eq_total_ref = f"Equipment!$F${r}"

    # -------------------------------------------------------- Subcontractors
    wsub = wb.create_sheet("Subcontractors")
    _title(wsub, est, "Subcontractors & allowances (R-006)", approved)
    _hdr(wsub, 4, ["Code", "Description", "Qty", "Unit", "Unit cost ($)", "Cost ($)", "Schedule of values",
                   "Basis"], [9, 36, 8, 8, 13, 13, 26, 60])
    r = 5
    subs = est["subcontractors"] or [{"code": "", "description": "(none - add subcontract quotes here)",
                                      "quantity": 0, "unit": "ls", "unit_cost": 0,
                                      "sov": "Other / added scope", "basis": ""}]
    for sb in subs:
        _cell(wsub, r, 1, sb["code"])
        _cell(wsub, r, 2, sb["description"])
        _cell(wsub, r, 3, sb["quantity"], F_IN, QTY)
        _cell(wsub, r, 4, sb["unit"])
        _cell(wsub, r, 5, sb["unit_cost"], F_IN, USD)
        _cell(wsub, r, 6, f"=ROUND(C{r}*E{r},2)", F, USD)
        _cell(wsub, r, 7, sb["sov"])
        _cell(wsub, r, 8, sb["basis"], align=WRAP)
        r += 1
    _cell(wsub, r, 5, "Subcontract total", FB, fill=FILL_TOT)
    _cell(wsub, r, 6, f"=SUM(F5:F{r - 1})", FB, USD, FILL_TOT)
    sub_total_ref, sub_last = f"Subcontractors!$F${r}", r - 1

    # -------------------------------------------------------------- Summary
    ws = wb.create_sheet("Summary", 0)
    _title(ws, est, f"Bid estimate  |  prepared {est['estimate_date']}", approved)
    ws.column_dimensions["A"].width = 34
    ws.column_dimensions["B"].width = 26
    ws.column_dimensions["C"].width = 70
    info = [("Status", "APPROVED FOR RELEASE" if approved else "DRAFT - PENDING REVIEW"),
            ("Review gate", est["review_status"]),
            ("Project type", p["project_type"]), ("Location", p["location"]),
            ("Site access", p["site_access"]), ("Target duration (days)", p["target_duration_days"])]
    if approved:
        info += [("Approved by", approval.get("approved_by")), ("Approved on", approval.get("approved_at"))]
    r = 4
    for k, v in info:
        _cell(ws, r, 1, k, FB)
        cc = _cell(ws, r, 2, v)
        if k == "Review gate" and v != "READY FOR REVIEW":
            cc.fill = FILL["CLARIFY"]
        r += 1
    r += 1
    _cell(ws, r, 1, "Cost build-up (R-006 to R-008)", F_HDR, fill=FILL_HDR)
    _cell(ws, r, 2, "Amount ($)", F_HDR, fill=FILL_HDR)
    _cell(ws, r, 3, "Basis", F_HDR, fill=FILL_HDR)
    r += 1
    top = r
    rows = [("Materials", f"={mat_total_ref}", "Materials sheet"),
            ("Labor", f"={lab_total_ref}", "Labor sheet (task hours + supervision + crew support time)"),
            ("Equipment rental", f"={eq_total_ref}", "Equipment sheet"),
            ("Subcontractors", f"={sub_total_ref}", "Subcontractors sheet")]
    for k, f, b in rows:
        _cell(ws, r, 1, k)
        _cell(ws, r, 2, f, F_LINK, USD)
        _cell(ws, r, 3, b)
        r += 1
    direct_r = r
    _cell(ws, r, 1, "Direct cost subtotal", FB, fill=FILL_TOT)
    _cell(ws, r, 2, f"=SUM(B{top}:B{r - 1})", FB, USD, FILL_TOT)
    _cell(ws, r, 3, "R-006: materials + labor + equipment + subcontractors", fill=FILL_TOT)
    r += 1
    cpct_r = r
    _cell(ws, r, 1, "Contingency %", FB)
    _cell(ws, r, 2, s["contingency_pct"], F_IN, PCT, FILL_KEY)
    _cell(ws, r, 3, s["contingency_basis"])
    r += 1
    cont_r = r
    _cell(ws, r, 1, "Contingency")
    _cell(ws, r, 2, f"=ROUND(B{direct_r}*B{cpct_r},2)", F, USD)
    _cell(ws, r, 3, "R-007: direct cost x contingency %")
    r += 1
    mpct_r = r
    _cell(ws, r, 1, "Overhead & profit %", FB)
    _cell(ws, r, 2, s["markup_pct"], F_IN, PCT, FILL_KEY)
    _cell(ws, r, 3, s["markup_basis"])
    r += 1
    mk_r = r
    _cell(ws, r, 1, "Overhead & profit")
    _cell(ws, r, 2, f"=ROUND((B{direct_r}+B{cont_r})*B{mpct_r},2)", F, USD)
    _cell(ws, r, 3, "R-008: (direct + contingency) x markup %")
    r += 1
    total_r = r
    _cell(ws, r, 1, "TOTAL BID", Font(name=FONT, size=12, bold=True), fill=FILL_TOT)
    _cell(ws, r, 2, f"=B{direct_r}+B{cont_r}+B{mk_r}", Font(name=FONT, size=12, bold=True), USD, FILL_TOT)
    _cell(ws, r, 3, "Direct + contingency + overhead & profit", fill=FILL_TOT)
    total_ref = f"Summary!$B${total_r}"
    r += 2

    st = est["comparables"].get("stats") or {}
    _cell(ws, r, 1, "Historical calibration (R-010)", F_HDR, fill=FILL_HDR)
    _cell(ws, r, 2, "Value", F_HDR, fill=FILL_HDR)
    _cell(ws, r, 3, "Note", F_HDR, fill=FILL_HDR)
    r += 1
    if st:
        med_r = r + 1
        for k, v, fmt, note in [
            ("Comparable final cost - low", st["final_cost_min"], USD0, f"{st['count']} closest same-type records"),
            ("Comparable final cost - median", st["final_cost_median"], USD0, "See Comparables sheet"),
            ("Comparable final cost - high", st["final_cost_max"], USD0, ""),
            ("Bid vs comparable median", f"=B{total_r}/B{med_r}-1", PCT, "Do not copy historical totals blindly"),
            ("Comparable median crew", st["crew_median"], "0.0", f"Estimate crew: {c['crew_size']}"),
            ("Comparable mean overrun (days)", st["overrun_days_mean"], "0.0", "R-012 schedule risk"),
        ]:
            _cell(ws, r, 1, k)
            _cell(ws, r, 2, v, F, fmt)
            _cell(ws, r, 3, note)
            r += 1
    r += 1
    n = {lv: sum(1 for f in est["flags"] if f["level"] == lv) for lv in ("CLARIFY", "REVIEW", "INFO")}
    _cell(ws, r, 1, "Review flags", FB)
    _cell(ws, r, 2, f"{n['CLARIFY']} clarify / {n['REVIEW']} review / {n['INFO']} info")
    _cell(ws, r, 3, "See 'Review Flags' sheet. CLARIFY items must be resolved before sending (R-009, R-013).")
    for row in ws.iter_rows(min_row=4, max_row=r, min_col=3, max_col=3):
        for cell in row:
            cell.alignment = WRAP

    # ---------------------------------------------------- Schedule of values
    wv = wb.create_sheet("Schedule of Values", 1)
    _title(wv, est, "Client schedule of values (markups distributed pro rata)", approved)
    _hdr(wv, 4, ["Work category", "Direct cost ($)", "Share", "Contingency + OH&P ($)", "Bid amount ($)"],
         [40, 16, 10, 20, 16])
    r = 5
    cats = list(est["schedule_of_values"].keys())
    gc = "General conditions, supervision & equipment"
    for cat in cats:
        _cell(wv, r, 1, cat)
        f = (f"=SUMIF(Materials!$J${mat_first}:$J${mat_last},A{r},Materials!$I${mat_first}:$I${mat_last})"
             f"+SUMIF(Labor!$J${lab_first}:$J${lab_last},A{r},Labor!$I${lab_first}:$I${lab_last})"
             f"+SUMIF(Subcontractors!$G$5:$G${sub_last},A{r},Subcontractors!$F$5:$F${sub_last})")
        if cat == gc:
            f += f"+{eq_total_ref}"
        _cell(wv, r, 2, f, F, USD)
        r += 1
    last = r - 1
    for i in range(5, r):
        _cell(wv, i, 3, f"=IF(SUM($B$5:$B${last})=0,0,B{i}/SUM($B$5:$B${last}))", F, PCT)
        _cell(wv, i, 4, f"=C{i}*(Summary!$B${cont_r}+Summary!$B${mk_r})", F, USD)
        _cell(wv, i, 5, f"=B{i}+D{i}", F, USD)
    _cell(wv, r, 1, "Total", FB, fill=FILL_TOT)
    for col, L in ((2, "B"), (4, "D"), (5, "E")):
        _cell(wv, r, col, f"=SUM({L}5:{L}{last})", FB, USD, FILL_TOT)
    _cell(wv, r, 3, None, fill=FILL_TOT)
    _cell(wv, r + 1, 1, "Check vs Summary total (should be 0)")
    _cell(wv, r + 1, 5, f"=ROUND(E{r}-{total_ref},2)", F, USD)

    # ---------------------------------------------------------- Review flags
    wf = wb.create_sheet("Review Flags", 2)
    _title(wf, est, "Review flags & clarifications (R-009 / R-012 / R-013)", approved)
    _hdr(wf, 4, ["Level", "Rule", "Message", "Resolved? (Y/N)", "Reviewer notes"], [10, 12, 90, 14, 40])
    dv = DataValidation(type="list", formula1='"Y,N"', allow_blank=True)
    wf.add_data_validation(dv)
    for i, fl in enumerate(est["flags"], start=5):
        _cell(wf, i, 1, fl["level"], FB, fill=FILL[fl["level"]])
        _cell(wf, i, 2, fl["rule"])
        _cell(wf, i, 3, fl["message"], align=WRAP)
        rc = _cell(wf, i, 4, "Y" if approved else "", F_IN, fill=FILL_KEY)
        dv.add(rc)
        _cell(wf, i, 5, None, F_IN)
    wf.cell(row=len(est["flags"]) + 6, column=1,
            value="CLARIFY = must be answered by the client before the bid is sent. REVIEW = estimator judgment. "
                  "INFO = assumption recorded for transparency.").font = Font(name=FONT, size=9, italic=True)

    # ----------------------------------------------------------- Comparables
    wc = wb.create_sheet("Comparables")
    _title(wc, est, "Historical comparables (R-010) - " + est["comparables"].get("method", ""), approved)
    comps = est["comparables"].get("comparables") or []
    if comps:
        keys = list(comps[0].keys())
        _hdr(wc, 4, keys, [12 if k != "notes" else 44 for k in keys])
        for i, comp in enumerate(comps, start=5):
            for j, k in enumerate(keys, start=1):
                fmt = USD0 if k == "final_cost_usd" else PCT if k.endswith("_pct") else None
                _cell(wc, i, j, comp[k], F, fmt)
    # ---------------------------------------------------------- Sample check
    wsc = wb.create_sheet("Sample Check")
    _title(wsc, est, "Cross-check against Sample_Bid_Lines", approved)
    _hdr(wsc, 4, ["Line item", "Unit", "Sample order qty", "Estimate order qty", "Difference"], [36, 9, 16, 16, 12])
    for i, sc in enumerate(est["sample_check"], start=5):
        _cell(wsc, i, 1, sc["line_item"])
        _cell(wsc, i, 2, sc["unit"])
        _cell(wsc, i, 3, sc["sample_qty"], F, QTY)
        _cell(wsc, i, 4, sc["estimate_qty"], F, QTY)
        _cell(wsc, i, 5, f"=D{i}-C{i}", F, QTY)

    # ---------------------------------------------------- Scope & assumptions
    wa = wb.create_sheet("Scope & Assumptions")
    _title(wa, est, "Project request, reviewer overrides and estimator assumptions", approved)
    wa.column_dimensions["A"].width = 34
    wa.column_dimensions["B"].width = 22
    wa.column_dimensions["C"].width = 22
    _hdr(wa, 4, ["Request field", "As requested", "As estimated"])
    r = 5
    for k, v in est["request"].items():
        _cell(wa, r, 1, k)
        _cell(wa, r, 2, v)
        est_v = est["project"].get(k)
        c3 = _cell(wa, r, 3, est_v)
        if str(est_v) != str(v) and not (isinstance(v, (int, float)) and float(v) == float(est_v or 0)):
            c3.fill = FILL["REVIEW"]
            c3.comment = Comment("Reviewer override", "Bid Agent")
        r += 1
    r += 1
    _cell(wa, r, 1, "Estimator assumption", F_HDR, fill=FILL_HDR)
    _cell(wa, r, 2, "Value", F_HDR, fill=FILL_HDR)
    _cell(wa, r, 3, "Source", F_HDR, fill=FILL_HDR)
    r += 1

    def walk(d, prefix=""):
        nonlocal r
        for k, v in d.items():
            if k.startswith("_"):
                continue
            if isinstance(v, dict):
                walk(v, f"{prefix}{k}.")
            else:
                _cell(wa, r, 1, prefix + k)
                _cell(wa, r, 2, ", ".join(map(str, v)) if isinstance(v, list) else v)
                _cell(wa, r, 3, "config/assumptions.json")
                r += 1
    walk(est["assumptions"])
    if est["overrides"]:
        r += 1
        _cell(wa, r, 1, "Reviewer overrides (JSON)", FB)
        import json
        _cell(wa, r, 2, json.dumps(est["overrides"]), align=WRAP)
        wa.merge_cells(start_row=r, start_column=2, end_row=r, end_column=3)
    r += 2
    _cell(wa, r, 1, "Standard exclusions", FB)
    for ex in est["exclusions"]:
        r += 1
        _cell(wa, r, 1, ex, border=False)

    for wsx in wb.worksheets:
        wsx.sheet_view.showGridLines = False
        wsx.page_setup.orientation = "landscape"
        wsx.page_setup.fitToWidth = 1
        wsx.sheet_properties.pageSetUpPr.fitToPage = True
        wsx.page_setup.fitToHeight = 0
    _steps_sheet(wb, est, approved)
    from openpyxl.workbook.properties import CalcProperties
    wb.calculation = CalcProperties(fullCalcOnLoad=True)
    wb.save(path)
    return path


def _steps_sheet(wb, est: dict, approved: bool):
    """Plain-language walkthrough of every calculation step, with the final answer first."""
    from .explain import build_steps, final_answer
    ws = wb.create_sheet("Calculation Steps", 1)
    _title(ws, est, "How this estimate was built, step by step", approved)
    ws.column_dimensions["A"].width = 8
    ws.column_dimensions["B"].width = 34
    ws.column_dimensions["C"].width = 110
    fa = final_answer(est)
    b = fa["breakdown"]
    r = 4
    _cell(ws, r, 1, "FINAL ANSWER", F_HDR, fill=FILL_HDR)
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=3)
    r += 1
    rows = [("Total bid", fa["total_bid"], USD), ("Materials", b["materials"], USD), ("Labor", b["labor"], USD),
            ("Equipment", b["equipment"], USD), ("Subcontractors", b["subcontractors"], USD),
            ("Direct cost", b["direct_cost"], USD),
            (f"Contingency ({b['contingency_pct']:.1%})", b["contingency"], USD),
            (f"Overhead & profit ({b['overhead_profit_pct']:.1%})", b["overhead_profit"], USD),
            ("Duration / crew", f"{fa['duration_days']} working days, crew of {fa['crew_size']}", None),
            ("Recommendation", fa["recommendation"], None)]
    rows += [("Client question", q, None) for q in fa["open_questions_for_client"]]
    rows += [("Reviewer item", q, None) for q in fa["items_for_reviewer"]]
    for k, v, fmt in rows:
        bold = k == "Total bid"
        _cell(ws, r, 2, k, FB)
        c = _cell(ws, r, 3, v, FB if bold else F, fmt, FILL_TOT if bold else None,
                  align=Alignment(wrap_text=True, vertical="top", horizontal="left"))
        if k == "Client question":
            c.fill = FILL["CLARIFY"]
        elif k == "Reviewer item":
            c.fill = FILL["REVIEW"]
        r += 1
    ws.cell(row=r, column=2, value="Values as generated by the agent. If you edit blue inputs on other sheets, "
            "the Summary sheet is the live total.").font = Font(name=FONT, size=9, italic=True)
    r += 2
    _hdr(ws, r, ["Step", "What was done (rule)", "Calculation and result"])
    ws.freeze_panes = None
    r += 1
    for stp in build_steps(est):
        _cell(ws, r, 1, stp["step"], FB, align=Alignment(horizontal="center", vertical="top"))
        _cell(ws, r, 2, f"{stp['title']}\n({stp['rule']})", FB, align=WRAP)
        text = "\n".join("- " + d for d in stp["details"]) + f"\n=> {stp['result']}"
        _cell(ws, r, 3, text, align=WRAP)
        ws.row_dimensions[r].height = min(409, 14 * (text.count("\n") + 1 + sum(len(d) // 120 for d in stp["details"])))
        r += 1
    ws.sheet_view.showGridLines = False
    ws.page_setup.orientation = "landscape"
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0
